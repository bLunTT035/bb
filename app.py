from flask import Flask, render_template_string, request, redirect
import os
from datetime import datetime
import re

app = Flask(__name__)

produtos_db = [
    {"id": 1, "nome": "Minoxidil 60ml", "categoria": "Minoxidil", "preco": 85.00},
    {"id": 2, "nome": "Pomada Modeladora", "categoria": "Pomada", "preco": 30.00},
    {"id": 3, "nome": "Gel Cola", "categoria": "Gel", "preco": 25.00},
    {"id": 4, "nome": "Creme de Barbear", "categoria": "Creme", "preco": 35.00},
    {"id": 5, "nome": "Shampoo 2 em 1", "categoria": "Shampoo", "preco": 28.00},
    {"id": 6, "nome": "Balm Pós-Barba", "categoria": "Balm", "preco": 32.00},
]
atendimentos_db = []

planos_mensal = [
    {"id": 1, "nome": "CORTE", "valor": 85},
    {"id": 2, "nome": "BARBA", "valor": 85},
    {"id": 3, "nome": "CORTE + SOBRANCELHA", "valor": 90},
    {"id": 4, "nome": "CORTE + BARBA E SOBRANCELHA", "valor": 125},
]

servicos_avulso = [
    {"id": 1, "nome": "Corte Simples", "valor": 35},
    {"id": 2, "nome": "Barba", "valor": 30},
    {"id": 3, "nome": "Corte + Barba", "valor": 60},
    {"id": 4, "nome": "Pezinho", "valor": 15},
    {"id": 5, "nome": "Sobrancelha", "valor": 10},
]

# ADICIONADO TELEFONE
mensalistas_db = [
    {"id": 1, "nome": "João Silva", "telefone": "(38) 99999-0000", "plano": "CORTE + BARBA E SOBRANCELHA", "valor": 125.0, "vencimento": 10, "status": "Pendente", "cortes": 2, "ultimo_pag": "01/10/2026", "historico": ["02/10 - Corte"]},
]

# NOVO - CONTROLE DE CLIENTES AVULSOS QUE VOLTARAM
clientes_avulso_db = []

def next_id(lista):
    return max([x["id"] for x in lista], default=0) + 1

HTML_BASE = """
<!DOCTYPE html>
<html>
<head>
<title>Rei da Navalha - PRO</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;700;900&display=swap" rel="stylesheet">
<style>
    *{font-family:'Inter',sans-serif; box-sizing:border-box}
    body { margin:0; background: linear-gradient(135deg, #0d2d6b 0%, #06163a 100%); min-height:100vh; }
 .header { background: rgba(255,255,255,0.97); border-bottom: 4px solid #c8102e; padding: 14px 24px; display:flex; justify-content:space-between; align-items:center; position:sticky; top:0; z-index:100; flex-wrap:wrap; gap:10px; }
 .header h1 { margin:0; color:#000; font-size:23px; font-weight:900; display:flex; align-items:center; gap:8px; flex:1; }
 .header h1 span { color:#c8102e; }
 .nav { display:flex; align-items:center; flex-wrap:wrap; gap:6px; }
 .nav a { text-decoration:none; color:#000; background:#f1f5f9; padding:8px 14px; border-radius:100px; font-size:13px; font-weight:800; border:1.5px solid #e2e8f0; white-space:nowrap; }
 .nav a.active { background:#0d2d6b; color:white; }
 .nav a.cta { background:#c8102e; color:white; }
 .hamburger { display:none; background:#0d2d6b; color:white; border:none; border-radius:10px; padding:8px 12px; font-size:20px; font-weight:900; cursor:pointer; }
 .container { max-width: 950px; margin: 28px auto; padding: 16px; }
 .card { background: #fff; border-radius: 20px; padding: 24px; box-shadow: 0 15px 35px rgba(0,0,0,0.25); position:relative; overflow:hidden; margin-bottom:16px; }
 .card::before{content:''; position:absolute; top:0; left:0; right:0; height:5px; background:linear-gradient(90deg, #0d2d6b, #c8102e);}
 .card h2 { color:#000!important; margin-top:0; font-weight:900; font-size:18px; display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:8px; }
 .grid3{ display:grid; grid-template-columns:1fr 1fr 1fr; gap:14px; margin-bottom:22px; }
 .stat { background:#fff; padding:18px; border-radius:18px; text-align:center; box-shadow:0 8px 20px rgba(0,0,0,0.2); }
 .stat b{ display:block; font-size:11px; text-transform:uppercase; color:#64748b; margin-bottom:6px; }
 .stat h3{ margin:0; font-size:26px; font-weight:900; }
    label { color: #000; font-weight: 800; font-size:12px; display:block; margin-top:14px; text-transform:uppercase; }
    input, select { width:100%; padding:12px 14px; border:2px solid #e2e8f0; border-radius:12px; margin-top:6px; box-sizing:border-box; color:#000; font-weight:700; background:#f8fafc; }
 .btn { background: linear-gradient(135deg, #c8102e 0%, #8f0c22 100%); color: white; border: none; padding: 14px; width:100%; border-radius:12px; font-weight:900; margin-top:18px; cursor:pointer; text-transform:uppercase; font-size:14px; }
 .btn-blue{ background: linear-gradient(135deg, #0d2d6b 0%, #1e4bb8 100%); }
 .btn-small{ padding:6px 10px; border-radius:8px; border:none; font-weight:800; font-size:11px; cursor:pointer; }
 .prod-item { border:2px solid #f1f5f9; padding:12px 14px; border-radius:14px; margin-bottom:10px; display:flex; justify-content:space-between; align-items:center; cursor:pointer; background:#f8fafc; color:#000; }
 .prod-item.selected { background:#0d2d6b; color:white; border-color:#0d2d6b; }
 .tag { background:#c8102e; color:white; font-size:10px; padding:3px 8px; border-radius:100px; font-weight:800; }
 .tag-green { background:#16a34a; }.tag-orange { background:#f59e0b; }
 .mini-btn{ padding:7px 12px; border-radius:8px; border:none; font-weight:800; font-size:12px; cursor:pointer; margin-left:6px; text-decoration:none; display:inline-block; }
 .badge-pago{ background:#dcfce7; color:#166534; padding:4px 10px; border-radius:100px; font-size:11px; font-weight:900; }
 .badge-pend{ background:#fee2e2; color:#991b1b; padding:4px 10px; border-radius:100px; font-size:11px; font-weight:900; }
 .valor-card{ background:#f8fafc; padding:12px; border-radius:12px; display:flex; justify-content:space-between; align-items:center; margin-bottom:8px; border-left:4px solid #0d2d6b; }
 .valor-card.mensal{ border-left-color:#c8102e; }
 .valor-actions{ display:flex; gap:4px; }
 .import-box{ background:#f8fafc; border:2px dashed #0d2d6b; padding:16px; border-radius:14px; margin-top:12px; }
    @media(max-width:768px){
   .header { padding:10px 12px; }.header h1 { font-size:18px; }.hamburger { display:block; }
   .nav { display:none; width:100%; flex-direction:column; background:#fff; border-radius:14px; padding:10px; box-shadow:0 8px 20px rgba(0,0,0,0.15); }
   .nav.show { display:flex; }.nav a { width:100%; text-align:center; padding:12px; }
   .grid3 { grid-template-columns:1fr; }.container { margin:14px auto; padding:10px; }
    }
</style>
</head>
<body>
<div class="header">
  <h1>👑 REI DA <span>NAVALHA</span></h1>
  <button class="hamburger" onclick="toggleMenu()">☰</button>
  <div class="nav" id="navMenu">
    <a href="/" class="{{'active' if active=='inicio' else ''}}">Início</a>
    <a href="/mensalistas" class="{{'active' if active=='mensal' else ''}}">💳 Mensalistas</a>
    <a href="/clientes" class="{{'active' if active=='clientes' else ''}}">👥 Avulsos</a>
    <a href="/produtos" class="{{'active' if active=='prod' else ''}}">Produtos</a>
    <a href="/historico" class="{{'active' if active=='hist' else ''}}">Histórico</a>
    <a href="/novo" class="cta">+ Novo</a>
  </div>
</div>
<div class="container">
{{content}}
</div>
<script>
function toggleMenu(){ document.getElementById('navMenu').classList.toggle('show'); }
function toggleProd(id){const el=document.getElementById('p-'+id); const inp=document.getElementById('produtos_input'); let s=inp.value?inp.value.split(',').filter(x=>x):[]; if(el.classList.contains('selected')){el.classList.remove('selected'); s=s.filter(x=>x!=id);}else{el.classList.add('selected'); s.push(id);} inp.value=s.join(','); calcTotal();}
function atualizaServico(){ const sel=document.getElementById('servico_select'); if(!sel) return; const preco=parseFloat(sel.options[sel.selectedIndex].dataset.preco)||0; document.getElementById('valor_servico').value=preco; calcTotal(); }
function atualizaPlanoMensal(){ const sel=document.getElementById('plano_mensal_select'); if(!sel) return; const preco=sel.options[sel.selectedIndex].dataset.preco; const inp=document.getElementById('valor_mensal_input'); if(inp) inp.value=preco; }
function calcTotal(){ const el=document.getElementById('valor_servico'); if(!el) return; const v=parseFloat(el.value)||0; let totP=0; const inp=document.getElementById('produtos_input'); if(inp) (inp.value.split(',').filter(x=>x)).forEach(id=>{const p=document.getElementById('preco-'+id); if(p) totP+=parseFloat(p.value);}); const t=v+totP; const ta=document.getElementById('total_auto'); const tv=document.getElementById('total_view'); if(ta) ta.value=t.toFixed(2); if(tv) tv.innerText='R$ '+t.toFixed(2); }
</script>
</body>
</html>
"""

@app.route("/", methods=["GET","POST"])
def index():
    global planos_mensal, servicos_avulso
    if request.method=="POST":
        acao = request.form.get("acao")
        if acao=="novo_plano":
            planos_mensal.append({"id": next_id(planos_mensal), "nome": request.form.get("nome").upper(), "valor": float(request.form.get("valor") or 0)})
        elif acao=="editar_plano":
            pid=int(request.form.get("id")); p=next((x for x in planos_mensal if x["id"]==pid),None)
            if p: p["nome"]=request.form.get("nome").upper(); p["valor"]=float(request.form.get("valor") or 0)
        elif acao=="excluir_plano":
            pid=int(request.form.get("id")); planos_mensal=[x for x in planos_mensal if x["id"]!=pid]
        elif acao=="novo_avulso":
            servicos_avulso.append({"id": next_id(servicos_avulso), "nome": request.form.get("nome"), "valor": float(request.form.get("valor") or 0)})
        elif acao=="editar_avulso":
            pid=int(request.form.get("id")); p=next((x for x in servicos_avulso if x["id"]==pid),None)
            if p: p["nome"]=request.form.get("nome"); p["valor"]=float(request.form.get("valor") or 0)
        elif acao=="excluir_avulso":
            pid=int(request.form.get("id")); servicos_avulso=[x for x in servicos_avulso if x["id"]!=pid]
        return redirect("/")

    total_vendas = sum(a['total'] for a in atendimentos_db)
    total_mensal = sum(m['valor'] for m in mensalistas_db if m['status']=='Pago')
    pend = len([m for m in mensalistas_db if m['status']=='Pendente'])
    edit_tipo = request.args.get("edit_tipo")
    edit_id = request.args.get("edit_id", type=int)

    html_planos=""
    for p in planos_mensal:
        if edit_tipo=="plano" and edit_id==p["id"]:
            html_planos+=f"""
            <form method="POST" style="background:#fff0f2; padding:10px; border-radius:12px; margin-bottom:8px; display:grid; grid-template-columns:1.5fr 0.5fr auto; gap:6px;">
              <input type="hidden" name="acao" value="editar_plano"><input type="hidden" name="id" value="{p['id']}">
              <input name="nome" value="{p['nome']}" required><input name="valor" type="number" value="{p['valor']}">
              <div><button class="btn-small" style="background:#0d2d6b; color:white;">Salvar</button><a href="/" class="btn-small" style="background:#e2e8f0; text-decoration:none; color:#000; padding:6px 8px;">X</a></div>
            </form>"""
        else:
            html_planos+=f"""<div class="valor-card mensal"><div><b style="font-size:12px;">{p['nome']}</b><br><span style="color:#c8102e; font-weight:900;">R$ {p['valor']:.0f}</span></div>
              <div class="valor-actions"><a href="/?edit_tipo=plano&edit_id={p['id']}" class="btn-small" style="background:#fef3c7; color:#92400e; text-decoration:none;">Editar</a>
              <form method="POST" style="display:inline;" onsubmit="return confirm('Excluir plano?')"><input type="hidden" name="acao" value="excluir_plano"><input type="hidden" name="id" value="{p['id']}"><button class="btn-small" style="background:#fee2e2; color:#991b1b;">X</button></form></div></div>"""

    html_avulso=""
    for s in servicos_avulso:
        if edit_tipo=="avulso" and edit_id==s["id"]:
            html_avulso+=f"""<form method="POST" style="background:#f1f5f9; padding:10px; border-radius:12px; margin-bottom:8px; display:grid; grid-template-columns:1.5fr 0.5fr auto; gap:6px;">
              <input type="hidden" name="acao" value="editar_avulso"><input type="hidden" name="id" value="{s['id']}">
              <input name="nome" value="{s['nome']}" required><input name="valor" type="number" value="{s['valor']}">
              <div><button class="btn-small" style="background:#0d2d6b; color:white;">Salvar</button><a href="/" class="btn-small" style="background:#e2e8f0; text-decoration:none; color:#000; padding:6px 8px;">X</a></div></form>"""
        else:
            html_avulso+=f"""<div class="valor-card"><div><b style="font-size:12px;">{s['nome']}</b><br><span style="color:#0d2d6b; font-weight:900;">R$ {s['valor']:.0f}</span></div>
              <div class="valor-actions"><a href="/?edit_tipo=avulso&edit_id={s['id']}" class="btn-small" style="background:#fef3c7; color:#92400e; text-decoration:none;">Editar</a>
              <form method="POST" style="display:inline;" onsubmit="return confirm('Excluir?')"><input type="hidden" name="acao" value="excluir_avulso"><input type="hidden" name="id" value="{s['id']}"><button class="btn-small" style="background:#fee2e2; color:#991b1b;">X</button></form></div></div>"""

    content = f"""
    <div class="grid3">
      <div class="stat"><b>Atendimentos</b><h3 style="color:#c8102e;">{len(atendimentos_db)}</h3></div>
      <div class="stat"><b>Faturamento Avulso</b><h3 style="color:#0d2d6b;">R$ {total_vendas:.2f}</h3></div>
      <div class="stat"><b>Mensalistas ({pend} pend.)</b><h3>R$ {total_mensal:.2f}</h3></div>
    </div>
    <div class="card">
      <h2>👑 Plano Mensal <button onclick="document.getElementById('form-plano').style.display='block'" class="btn-small" style="background:#c8102e; color:white;">+ Add</button></h2>
      <div id="form-plano" style="display:none; background:#f8fafc; padding:12px; border-radius:12px; margin-bottom:12px;">
        <form method="POST" style="display:grid; grid-template-columns:1.5fr 0.7fr auto; gap:8px;">
          <input type="hidden" name="acao" value="novo_plano"><input name="nome" placeholder="Ex: CORTE" required><input name="valor" type="number" placeholder="R$" required>
          <button class="btn-small" style="background:#0d2d6b; color:white; padding:12px;">Adicionar</button>
        </form>
      </div>
      <div style="display:grid; grid-template-columns:1fr 1fr; gap:8px;">{html_planos}</div>
    </div>
    <div class="card">
      <h2>✂️ Avulso <button onclick="document.getElementById('form-avulso').style.display='block'" class="btn-small" style="background:#0d2d6b; color:white;">+ Add</button></h2>
      <div id="form-avulso" style="display:none; background:#f8fafc; padding:12px; border-radius:12px; margin-bottom:12px;">
        <form method="POST" style="display:grid; grid-template-columns:1.5fr 0.7fr auto; gap:8px;">
          <input type="hidden" name="acao" value="novo_avulso"><input name="nome" placeholder="Ex: Corte + Barba" required><input name="valor" type="number" placeholder="R$" required>
          <button class="btn-small" style="background:#c8102e; color:white; padding:12px;">Adicionar</button>
        </form>
      </div>
      <div style="display:grid; grid-template-columns:1fr 1fr; gap:8px;">{html_avulso}</div>
      <div style="display:grid; grid-template-columns:1fr 1fr; gap:12px; margin-top:18px;">
        <a href="/novo"><button class="btn btn-blue">+ Novo Atendimento</button></a>
        <a href="/clientes"><button class="btn">👥 Ver Avulsos Frequentes</button></a>
      </div>
    </div>
    """
    return render_template_string(HTML_BASE.replace("{{content}}", content), active="inicio")

@app.route("/produtos", methods=["GET","POST"])
def produtos_route():
    global produtos_db
    msg = ""
    if request.method=="POST":
        if 'arquivo' in request.files and request.files['arquivo'].filename!= '':
            file = request.files['arquivo']
            try:
                try:
                    import pandas as pd
                    df = pd.read_excel(file)
                    cols = {c.lower().strip(): c for c in df.columns}
                    def find_col(names):
                        for n in names:
                            for k,v in cols.items():
                                if n in k:
                                    return v
                        return None
                    col_nome = find_col(['produto','nome','descricao','descrição','item'])
                    col_preco = find_col(['valor revenda','preco','preço','valor','price','revenda'])
                    col_cat = find_col(['categoria','cat','tipo','grupo'])
                    count=0
                    for _, row in df.iterrows():
                        nome = str(row[col_nome]).strip() if col_nome and str(row[col_nome])!='nan' else ""
                        if not nome or nome.lower() in ['nan','none','']: continue
                        if any(p['nome'].lower()==nome.lower() for p in produtos_db): continue
                        try: preco = float(row[col_preco]) if col_preco else 0.0
                        except: preco = 0.0
                        cat = str(row[col_cat]).strip() if col_cat and str(row[col_cat])!='nan' else "Geral"
                        if cat.lower()=='nan' or not cat: cat="Geral"
                        if preco<=0: continue
                        produtos_db.append({"id": next_id(produtos_db), "nome": nome, "categoria": cat, "preco": preco})
                        count+=1
                    msg = f"<div style='background:#dcfce7; color:#166534; padding:12px; border-radius:10px; margin-bottom:12px; font-weight:800;'>✅ {count} produtos importados! {len(df)} linhas lidas.</div>"
                except ImportError:
                    import openpyxl
                    wb = openpyxl.load_workbook(file)
                    ws = wb.active
                    headers = [str(ws.cell(1,c).value).lower() if ws.cell(1,c).value else "" for c in range(1, ws.max_column+1)]
                    def idx_of(names):
                        for i,h in enumerate(headers):
                            for n in names:
                                if n in h: return i
                        return None
                    i_nome = idx_of(['produto','nome','descricao'])
                    i_preco = idx_of(['valor revenda','preco','valor'])
                    i_cat = idx_of(['categoria','tipo'])
                    count=0
                    for r in range(2, ws.max_row+1):
                        nome = str(ws.cell(r, (i_nome+1) if i_nome is not None else 1).value or "").strip()
                        if not nome or nome.lower()=='none': continue
                        if any(p['nome'].lower()==nome.lower() for p in produtos_db): continue
                        try: preco = float(ws.cell(r, (i_preco+1) if i_preco is not None else 2).value or 0)
                        except: preco=0
                        if preco<=0: continue
                        cat = str(ws.cell(r, (i_cat+1) if i_cat is not None else 3).value or "Geral").strip()
                        produtos_db.append({"id": next_id(produtos_db), "nome": nome, "categoria": cat, "preco": preco})
                        count+=1
                    msg = f"<div style='background:#dcfce7; color:#166534; padding:12px; border-radius:10px; margin-bottom:12px; font-weight:800;'>✅ {count} produtos importados!</div>"
            except Exception as e:
                msg = f"<div style='background:#fee2e2; color:#991b1b; padding:12px; border-radius:10px; margin-bottom:12px; font-weight:800;'>❌ Erro: {e}</div>"
        else:
            acao=request.form.get("acao")
            mid=request.form.get("id")
            if acao=="novo":
                produtos_db.append({"id": next_id(produtos_db),"nome": request.form.get("nome"),"categoria": request.form.get("categoria"),"preco": float(request.form.get("preco") or 0)})
            elif acao=="excluir" and mid:
                produtos_db = [x for x in produtos_db if str(x["id"])!=str(mid)]
            elif acao=="editar" and mid:
                p=next((x for x in produtos_db if str(x["id"])==str(mid)),None)
                if p:
                    p["nome"]=request.form.get("nome")
                    p["categoria"]=request.form.get("categoria")
                    p["preco"]=float(request.form.get("preco"))
            return redirect("/produtos")
    edit_id = request.args.get("edit", type=int)
    rows=""
    for p in produtos_db:
        if edit_id==p["id"]:
            rows+=f"""<tr style="background:#f1f5f9;"><td colspan="3" style="padding:12px;">
              <form method="POST" style="display:grid; grid-template-columns:2fr 1fr 1fr 1fr; gap:8px;">
                <input type="hidden" name="acao" value="editar"><input type="hidden" name="id" value="{p['id']}">
                <input name="nome" value="{p['nome']}" required><input name="categoria" value="{p['categoria']}"><input name="preco" type="number" step="0.01" value="{p['preco']}">
                <div><button class="mini-btn" style="background:#0d2d6b; color:white;">Salvar</button><a href="/produtos" class="mini-btn" style="background:#e2e8f0;">Cancelar</a></div>
              </form></td></tr>"""
        else:
            rows+=f"""<tr style='color:#000;'><td style='padding:12px; border-bottom:1px solid #e2e8f0;'><span class='tag'>{p['categoria']}</span> <b style='margin-left:6px;'>{p['nome']}</b></td>
            <td style='padding:12px; border-bottom:1px solid #e2e8f0; font-weight:900;'>R$ {p['preco']:.2f}</td>
            <td style='padding:12px; border-bottom:1px solid #e2e8f0; text-align:right;'><a href="/produtos?edit={p['id']}" class="mini-btn" style="background:#fef3c7;">Editar</a>
            <form method="POST" style="display:inline;" onsubmit="return confirm('Excluir?')"><input type="hidden" name="acao" value="excluir"><input type="hidden" name="id" value="{p['id']}"><button class="mini-btn" style="background:#fee2e2;">Excluir</button></form></td></tr>"""
    content = f"""
    {msg}
    <div class="card">
      <h2>Novo Produto</h2>
      <form method="POST" style="display:grid; grid-template-columns:2fr 1fr 1fr 1fr; gap:10px; align-items:end;">
        <input type="hidden" name="acao" value="novo">
        <div><label>Nome</label><input name="nome" required></div>
        <div><label>Categoria</label><input name="categoria" value="Novo"></div>
        <div><label>Preco</label><input name="preco" type="number" step="0.01" required></div>
        <div><button class="btn btn-blue" style="margin-top:6px;">Add</button></div>
      </form>
      <div class="import-box">
        <h3 style="margin:0 0 8px 0; font-size:14px; font-weight:900; color:#0d2d6b;">📥 Importar Planilha (Excel)</h3>
        <p style="margin:0 0 10px 0; font-size:12px; color:#64748b;">Aceita.xlsx com colunas PRODUTO / VALOR REVENDA / CATEGORIA</p>
        <form method="POST" enctype="multipart/form-data" style="display:grid; grid-template-columns: 1fr auto; gap:10px; align-items:end;">
          <div><input type="file" name="arquivo" accept=".xlsx,.xls,.csv" required style="background:white;"></div>
          <div><button class="btn btn-blue" style="margin:0; padding:12px 20px;">IMPORTAR</button></div>
        </form>
      </div>
    </div>
    <div class="card">
      <h2>Produtos ({len(produtos_db)})</h2>
      <table style="width:100%; border-collapse:collapse; margin-top:12px;">
        <tr style="background:#0d2d6b; color:white;"><th style="padding:12px; text-align:left;">Produto</th><th>Preco</th><th>Acoes</th></tr>
        {rows}
      </table>
    </div>
    """
    return render_template_string(HTML_BASE.replace("{{content}}", content), active="prod")

@app.route("/novo", methods=["GET", "POST"])
def novo():
    global clientes_avulso_db
    if request.method == "POST":
        cliente = request.form.get("cliente")
        telefone = request.form.get("telefone")
        servico = request.form.get("servico")
        valor_servico = float(request.form.get("valor_servico") or 0)
        produtos_ids = request.form.get("produtos", "")
        total = float(request.form.get("total_auto") or 0)
        pagamento = request.form.get("pagamento")
        prod_nomes = []
        if produtos_ids:
            for pid in produtos_ids.split(","):
                if pid.strip().isdigit():
                    p = next((x for x in produtos_db if x["id"]==int(pid)), None)
                    if p: prod_nomes.append(p["nome"])
        data_hoje = datetime.now().strftime("%d/%m/%Y %H:%M")
        dia_hoje = datetime.now().strftime("%d/%m/%Y")
        atendimentos_db.append({"id": len(atendimentos_db)+1,"cliente": cliente,"telefone": telefone,"servico": servico,"valor_servico": valor_servico,"produtos": ", ".join(prod_nomes),"total": total,"pagamento": pagamento,"data": data_hoje})

        # ATUALIZA CONTROLE DE AVULSOS
        cli = next((x for x in clientes_avulso_db if x["nome"].lower()==cliente.lower()), None)
        if not cli:
            clientes_avulso_db.append({"id": next_id(clientes_avulso_db), "nome": cliente, "telefone": telefone, "primeira_visita": dia_hoje, "ultima_visita": dia_hoje, "visitas_mes": 1, "historico": [f"{dia_hoje} - {servico}"]})
        else:
            cli["telefone"]=telefone or cli["telefone"]
            cli["ultima_visita"]=dia_hoje
            cli["visitas_mes"]+=1
            cli["historico"].append(f"{dia_hoje} - {servico}")

        return redirect("/clientes")

    produtos_html = ""
    for p in produtos_db:
        produtos_html += f'<div class="prod-item" id="p-{p["id"]}" onclick="toggleProd({p["id"]})"><div><span class="tag">{p["categoria"]}</span> <b style="margin-left:6px;">{p["nome"]}</b><br><small>R$ {p["preco"]:.2f}</small></div><div>OK</div><input type="hidden" id="preco-{p["id"]}" value="{p["preco"]}"></div>'
    servicos_opt = "".join([f'<option value="{s["nome"]}" data-preco="{s["valor"]}">{s["nome"]} - R$ {s["valor"]:.0f}</option>' for s in servicos_avulso])
    content = f"""
    <div class="card"><h2>Novo Atendimento - Com Telefone</h2>
      <form method="POST">
        <div style="display:grid; grid-template-columns:1fr 1fr; gap:10px;">
          <div><label>Cliente *</label><input name="cliente" required></div>
          <div><label>Telefone *</label><input name="telefone" placeholder="(38) 99999-9999" required></div>
        </div>
        <label>Pagamento</label><select name="pagamento"><option>Dinheiro</option><option>Pix</option><option>Cartao</option><option>Mensalista</option></select>
        <div style="display:grid; grid-template-columns: 2fr 1fr 1fr; gap:10px;">
          <div><label>Servico Avulso</label><select name="servico" id="servico_select" onchange="atualizaServico()">{servicos_opt}</select></div>
          <div><label>Valor</label><input id="valor_servico" name="valor_servico" type="number" value="{servicos_avulso[0]['valor'] if servicos_avulso else 35}" oninput="calcTotal()"></div>
          <div><label>Total</label><input id="total_auto" name="total_auto" readonly style="font-weight:900; background:#fef3c7;" value="{servicos_avulso[0]['valor'] if servicos_avulso else 35}"><small id="total_view" style="color:#c8102e; font-weight:900;">R$ {servicos_avulso[0]['valor'] if servicos_avulso else 35}</small></div>
        </div>
        <label>Produtos - clique</label><input type="hidden" name="produtos" id="produtos_input"><div style="max-height:320px; overflow:auto; margin-top:8px;">{produtos_html}</div>
        <button class="btn" type="submit">SALVAR</button>
      </form>
    </div>"""
    return render_template_string(HTML_BASE.replace("{{content}}", content), active="novo")

@app.route("/mensalistas", methods=["GET","POST"])
def mensalistas():
    global mensalistas_db
    if request.method=="POST":
        acao = request.form.get("acao"); mid = request.form.get("id")
        if acao=="novo":
            plano_nome = request.form.get("plano")
            valor_plano = next((v["valor"] for v in planos_mensal if v["nome"]==plano_nome), 80)
            try: valor_plano = float(request.form.get("valor") or valor_plano)
            except: pass
            mensalistas_db.append({"id": next_id(mensalistas_db),"nome": request.form.get("nome"),"telefone": request.form.get("telefone"),"plano": plano_nome, "valor": float(valor_plano),"vencimento": int(request.form.get("vencimento") or 10),"status": "Pendente","cortes": 0,"ultimo_pag": "-","historico": []})
        elif acao=="pagar" and mid:
            m=next((x for x in mensalistas_db if str(x["id"])==str(mid)),None)
            if m: m["status"]="Pago"; m["ultimo_pag"]=datetime.now().strftime("%d/%m/%Y")
        elif acao=="cortar" and mid:
            m=next((x for x in mensalistas_db if str(x["id"])==str(mid)),None)
            if m: m["cortes"]+=1; m["historico"].append(datetime.now().strftime("%d/%m - Corte"))
        elif acao=="excluir" and mid:
            mensalistas_db = [x for x in mensalistas_db if str(x["id"])!=str(mid)]
        elif acao=="editar" and mid:
            m=next((x for x in mensalistas_db if str(x["id"])==str(mid)),None)
            if m: m["nome"]=request.form.get("nome"); m["telefone"]=request.form.get("telefone"); m["plano"]=request.form.get("plano"); m["valor"]=float(request.form.get("valor")); m["vencimento"]=int(request.form.get("vencimento"))
        return redirect("/mensalistas")
    edit_id = request.args.get("edit", type=int)
    options_planos = "".join([f'<option value="{p["nome"]}" data-preco="{p["valor"]}">{p["nome"]} - R$ {p["valor"]:.0f}</option>' for p in planos_mensal])
    lista=""
    for m in reversed(mensalistas_db):
        badge = "<span class='badge-pago'>PAGO</span>" if m["status"]=="Pago" else "<span class='badge-pend'>PENDENTE</span>"
        hist = "<br>".join(m["historico"][-3:]) if m.get("historico") else "<small style='color:#94a3b8;'>Nenhum</small>"
        form_edit = f"""<form method="POST" style="background:#f1f5f9; padding:12px; border-radius:10px; margin-top:12px;"><input type="hidden" name="acao" value="editar"><input type="hidden" name="id" value="{m['id']}"><div style="display:grid; grid-template-columns:2fr 1.5fr 1fr 1fr 1fr; gap:8px;"><input name="nome" value="{m['nome']}" required><input name="telefone" value="{m.get('telefone','')}" placeholder="Telefone"><select name="plano">{options_planos}</select><input name="valor" type="number" value="{m['valor']}"><input name="vencimento" type="number" value="{m['vencimento']}"></div><button class="mini-btn" style="background:#0d2d6b; color:white; margin-top:8px;">Salvar</button><a href="/mensalistas" class="mini-btn" style="background:#e2e8f0;">Cancelar</a></form>""" if edit_id==m["id"] else ""
        lista+=f"""<div class="card" style="border-left:5px solid #0d2d6b;"><div style="display:flex; justify-content:space-between;"><div><b>{m['nome']}</b> - {m.get('telefone','')}<br><small>{m.get('plano','')} - R$ {m['valor']:.2f}</small></div><div>{badge}</div></div><div style="margin-top:8px; background:#f8fafc; padding:8px; border-radius:8px; font-size:12px;">{hist}</div>{form_edit}
          <div style="margin-top:12px;"><form method="POST" style="display:inline;"><input type="hidden" name="acao" value="cortar"><input type="hidden" name="id" value="{m['id']}"><button class="mini-btn" style="background:#0d2d6b; color:white;">Corte</button></form>
          <form method="POST" style="display:inline;"><input type="hidden" name="acao" value="pagar"><input type="hidden" name="id" value="{m['id']}"><button class="mini-btn" style="background:#16a34a; color:white;">Pago</button></form>
          <a href="/mensalistas?edit={m['id']}" class="mini-btn" style="background:#fef3c7;">Editar</a>
          <form method="POST" style="display:inline;" onsubmit="return confirm('Excluir?')"><input type="hidden" name="acao" value="excluir"><input type="hidden" name="id" value="{m['id']}"><button class="mini-btn" style="background:#fee2e2;">Excluir</button></form></div></div>"""
    content=f"""<div class="card"><h2>Novo Mensalista</h2><form method="POST" style="display:grid; grid-template-columns:2fr 1.5fr 2fr 1fr 1fr 1fr; gap:10px; align-items:end;"><input type="hidden" name="acao" value="novo"><div><label>Nome</label><input name="nome" required></div><div><label>Telefone</label><input name="telefone" placeholder="(38) 9...." required></div><div><label>Plano</label><select name="plano" id="plano_mensal_select" onchange="atualizaPlanoMensal()">{options_planos}</select></div><div><label>Valor</label><input name="valor" id="valor_mensal_input" type="number" value="{planos_mensal[0]['valor'] if planos_mensal else 85}"></div><div><label>Venc.</label><input name="vencimento" type="number" value="10"></div><div><button class="btn btn-blue" style="margin-top:6px;">Add</button></div></form></div>{lista}"""
    return render_template_string(HTML_BASE.replace("{{content}}", content), active="mensal")

@app.route("/clientes", methods=["GET","POST"])
def clientes():
    global clientes_avulso_db
    if request.method=="POST":
        acao=request.form.get("acao"); mid=request.form.get("id")
        if acao=="novo":
            clientes_avulso_db.append({"id": next_id(clientes_avulso_db), "nome": request.form.get("nome"), "telefone": request.form.get("telefone"), "primeira_visita": request.form.get("primeira") or datetime.now().strftime("%d/%m/%Y"), "ultima_visita": request.form.get("ultima") or datetime.now().strftime("%d/%m/%Y"), "visitas_mes": int(request.form.get("visitas") or 1), "historico": [f"{datetime.now().strftime('%d/%m')} - Cadastro"]})
        elif acao=="editar" and mid:
            c=next((x for x in clientes_avulso_db if str(x["id"])==str(mid)), None)
            if c:
                c["nome"]=request.form.get("nome"); c["telefone"]=request.form.get("telefone"); c["primeira_visita"]=request.form.get("primeira"); c["ultima_visita"]=request.form.get("ultima"); c["visitas_mes"]=int(request.form.get("visitas") or 1)
        elif acao=="excluir" and mid:
            clientes_avulso_db=[x for x in clientes_avulso_db if str(x["id"])!=str(mid)]
        elif acao=="nova_visita" and mid:
            c=next((x for x in clientes_avulso_db if str(x["id"])==str(mid)), None)
            if c:
                hoje=datetime.now().strftime("%d/%m/%Y")
                c["ultima_visita"]=hoje; c["visitas_mes"]+=1; c["historico"].append(f"{hoje} - Visita avulsa")
        return redirect("/clientes")

    edit_id=request.args.get("edit", type=int)
    lista=""
    for c in reversed(clientes_avulso_db):
        tag = f'<span class="tag tag-green">{c["visitas_mes"]}x no mês</span>' if c["visitas_mes"]>=2 else f'<span class="tag">{c["visitas_mes"]}x</span>'
        alerta = '<span class="tag tag-orange">🔥 Virar mensalista?</span>' if c["visitas_mes"]>=2 else ''
        hist="<br>".join(c["historico"][-4:])
        form_edit=f"""
        <form method="POST" style="background:#f1f5f9; padding:12px; border-radius:10px; margin-top:10px; display:grid; grid-template-columns:1fr 1fr 1fr 1fr 0.5fr; gap:6px;">
          <input type="hidden" name="acao" value="editar"><input type="hidden" name="id" value="{c['id']}">
          <input name="nome" value="{c['nome']}"><input name="telefone" value="{c['telefone']}"><input name="primeira" value="{c['primeira_visita']}" placeholder="Dia veio"><input name="ultima" value="{c['ultima_visita']}" placeholder="Dia voltou"><input name="visitas" type="number" value="{c['visitas_mes']}">
          <div style="grid-column:1/-1;"><button class="mini-btn" style="background:#0d2d6b; color:white;">Salvar</button><a href="/clientes" class="mini-btn" style="background:#e2e8f0;">Cancelar</a></div>
        </form>""" if edit_id==c["id"] else ""
        lista+=f"""
        <div class="card" style="border-left:5px solid #0d2d6b;">
          <div style="display:flex; justify-content:space-between;"><div><b>{c['nome']}</b> - {c['telefone']}<br><small>Veio: {c['primeira_visita']} | Voltou: {c['ultima_visita']}</small></div><div>{tag} {alerta}</div></div>
          <div style="margin-top:8px; background:#f8fafc; padding:8px; border-radius:8px; font-size:12px;">{hist}</div>
          {form_edit}
          <div style="margin-top:10px;">
            <form method="POST" style="display:inline;"><input type="hidden" name="acao" value="nova_visita"><input type="hidden" name="id" value="{c['id']}"><button class="mini-btn" style="background:#0d2d6b; color:white;">+ Nova Visita</button></form>
            <a href="/clientes?edit={c['id']}" class="mini-btn" style="background:#fef3c7;">Editar</a>
            <form method="POST" style="display:inline;" onsubmit="return confirm('Excluir?')"><input type="hidden" name="acao" value="excluir"><input type="hidden" name="id" value="{c['id']}"><button class="mini-btn" style="background:#fee2e2;">Excluir</button></form>
          </div>
        </div>"""

    content=f"""
    <div class="card">
      <h2>👥 Controle Avulso - Quem veio mais de 1x no mês</h2>
      <form method="POST" style="display:grid; grid-template-columns:1.5fr 1fr 1fr 1fr 0.5fr 1fr; gap:8px; align-items:end;">
        <input type="hidden" name="acao" value="novo">
        <div><label>Nome</label><input name="nome" required></div>
        <div><label>Telefone</label><input name="telefone" required></div>
        <div><label>Dia Veio</label><input name="primeira" value="{datetime.now().strftime('%d/%m/%Y')}"></div>
        <div><label>Dia Voltou</label><input name="ultima" value="{datetime.now().strftime('%d/%m/%Y')}"></div>
        <div><label>Qtd</label><input name="visitas" type="number" value="1"></div>
        <div><button class="btn btn-blue">Add Cliente</button></div>
      </form>
    </div>
    {lista if lista else "<div class='card'><p>Nenhum cliente avulso ainda. Quando você lançar um atendimento no + Novo, ele já cria aqui automaticamente.</p></div>"}
    """
    return render_template_string(HTML_BASE.replace("{{content}}", content), active="clientes")

@app.route("/historico")
def historico():
    rows = ""
    for a in reversed(atendimentos_db):
        rows += f"<tr style='color:#000;'><td style='padding:10px; border-bottom:1px solid #e2e8f0;'>{a['data']}</td><td style='padding:10px; border-bottom:1px solid #e2e8f0; font-weight:800;'>{a['cliente']}<br><small>{a.get('telefone','')}</small></td><td style='padding:10px; border-bottom:1px solid #e2e8f0;'>{a['servico']}<br><small style='color:#c8102e;'>{a['produtos']}</small></td><td style='padding:10px; border-bottom:1px solid #e2e8f0; font-weight:900;'>R$ {a['total']:.2f}</td></tr>"
    if not rows: rows = "<tr><td colspan=4 style='padding:24px; text-align:center;'>Nenhum ainda</td></tr>"
    content = f"""<div class="card"><h2>Historico</h2><table style="width:100%; border-collapse:collapse; margin-top:12px;"><tr style="background:#0d2d6b; color:white;"><th style="padding:10px;">Data</th><th>Cliente</th><th>Servico</th><th>Total</th></tr>{rows}</table></div>"""
    return render_template_string(HTML_BASE.replace("{{content}}", content), active="hist")

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)
